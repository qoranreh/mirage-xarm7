#!/usr/bin/env python3
"""Build the Korean presentation for the current Panda -> xArm7 experiment."""

from __future__ import annotations

import json
import math
from pathlib import Path

import cv2
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from PIL import Image, ImageEnhance, ImageFilter, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "presentations"
ASSET_DIR = OUT_DIR / "assets"
PPTX_PATH = OUT_DIR / "mirage_xarm7_experiment_setup_2026-08-10.pptx"

RESULT_ROOT = ROOT / "mirage/mirage/benchmark/robosuite/results/xarm_transfer"
VIDEO_PATH = RESULT_ROOT / "panda_xarm_oracle_seed1_full_task.mp4"
FD_REPORT_PATH = RESULT_ROOT / "forward_dynamics_evaluation.json"
TRACK_ORACLE_PATH = RESULT_ROOT / "xarm_rollout_report_task_tool_v2.json"
TRACK_RIDGE_PATH = RESULT_ROOT / "xarm_learned_rollout_report_task_tool_v2.json"
PICK_REPORT_PATH = RESULT_ROOT / "pick-place-can/low_dim/xarm_rollout_report.json"
LIFT_REPORT_PATH = RESULT_ROOT / "lift/low_dim/xarm_rollout_report.json"

FONT = "Noto Sans CJK KR"
MONO = "Noto Sans Mono CJK KR"
FONT_FILE = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"

BG = "0B1020"
PANEL = "141B2D"
PANEL_2 = "1A2338"
WHITE = "F5F7FB"
TEXT = "DDE5F0"
MUTED = "96A3B7"
GRID = "2C3851"
GREEN = "39D98A"
GREEN_DARK = "173F35"
ORANGE = "FF914D"
ORANGE_DARK = "4A2B22"
PURPLE = "8B7CFF"
PURPLE_DARK = "29264A"
BLUE = "5BB7FF"
BLUE_DARK = "17354A"
RED = "FF6474"
YELLOW = "FFD166"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def load_json(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def set_fill(shape, color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)


def set_line(shape, color: str, width: float = 1.0) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = PANEL,
    line: str | None = None,
    radius: bool = True,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, fill)
    if line:
        set_line(shape, line)
    else:
        shape.line.fill.background()
    return shape


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = TEXT,
    bold: bool = False,
    font: str = FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.02,
    line_spacing: float | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    if line_spacing is not None:
        paragraph.line_spacing = line_spacing
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(
    slide,
    runs: list[tuple[str, str, bool]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    font: str = FONT,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = align
    for value, color, bold in runs:
        run = p.add_run()
        run.text = value
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_multiline(
    slide,
    lines: list[tuple[str, str, bool]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 16,
    bullet: bool = False,
    spacing_after: float = 6,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for idx, (value, color, bold) in enumerate(lines):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = ("•  " if bullet else "") + value
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = rgb(color)
        paragraph.space_after = Pt(spacing_after)
        paragraph.line_spacing = 1.04
    return box


def add_badge(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    *,
    fill: str = PANEL_2,
    color: str = TEXT,
    line: str | None = None,
    size: float = 11,
):
    shape = add_rect(slide, x, y, w, 0.34, fill=fill, line=line, radius=True)
    add_text(
        slide,
        text,
        x + 0.05,
        y + 0.005,
        w - 0.10,
        0.31,
        size=size,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    return shape


def add_title(slide, title: str, subtitle: str | None, number: int) -> None:
    add_text(slide, title, 0.62, 0.34, 10.8, 0.58, size=27, color=WHITE, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.64, 0.94, 11.5, 0.34, size=11.5, color=MUTED)
    add_text(
        slide,
        f"{number:02d}",
        11.92,
        0.38,
        0.76,
        0.34,
        size=12,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.64), Inches(1.27), Inches(12.04), Inches(0.018)
    )
    set_fill(line, GRID)
    line.line.fill.background()


def add_footer(slide, source: str = "") -> None:
    add_text(
        slide,
        "Local working tree · 2026-08-10 · SIMULATION ONLY",
        0.64,
        7.18,
        6.0,
        0.20,
        size=8.2,
        color="6F7D93",
    )
    if source:
        add_text(
            slide,
            source,
            6.15,
            7.18,
            6.55,
            0.20,
            size=8.0,
            color="6F7D93",
            align=PP_ALIGN.RIGHT,
        )


def add_picture_crop(slide, path: Path, x: float, y: float, w: float, h: float, border: str | None = None):
    with Image.open(path) as image:
        source_ratio = image.width / image.height
    box_ratio = w / h
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if source_ratio > box_ratio:
        crop = (1.0 - box_ratio / source_ratio) / 2.0
        picture.crop_left = crop
        picture.crop_right = crop
    elif source_ratio < box_ratio:
        crop = (1.0 - source_ratio / box_ratio) / 2.0
        picture.crop_top = crop
        picture.crop_bottom = crop
    if border:
        outline = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        outline.fill.background()
        set_line(outline, border, 1.1)
    return picture


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color: str = MUTED, width: float = 2.0):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    set_line(connector, color, width)
    connector.line.end_arrowhead = True
    return connector


def add_metric_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    value: str,
    label: str,
    detail: str,
    *,
    accent: str,
):
    add_rect(slide, x, y, w, h, fill=PANEL, line=GRID)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h)
    )
    set_fill(bar, accent)
    bar.line.fill.background()
    add_text(slide, value, x + 0.28, y + 0.20, w - 0.42, 0.62, size=28, color=accent, bold=True)
    add_text(slide, label, x + 0.30, y + 0.90, w - 0.45, 0.36, size=14, color=WHITE, bold=True)
    add_text(slide, detail, x + 0.30, y + 1.32, w - 0.45, h - 1.48, size=10.5, color=MUTED)


def make_video_assets() -> dict[int, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    capture = cv2.VideoCapture(str(VIDEO_PATH))
    frame_count = int(capture.get(cv2.CAP_PROP_FRAME_COUNT))
    if frame_count < 116:
        raise RuntimeError(f"unexpected video frame count: {frame_count}")
    selected = [0, 30, 60, 90, 115]
    paths: dict[int, Path] = {}
    for frame_id in selected:
        capture.set(cv2.CAP_PROP_POS_FRAMES, frame_id)
        ok, frame = capture.read()
        if not ok:
            raise RuntimeError(f"could not read video frame {frame_id}")
        output = ASSET_DIR / f"pick_place_step_{frame_id + 1:03d}.jpg"
        cv2.imwrite(str(output), frame, [int(cv2.IMWRITE_JPEG_QUALITY), 95])
        paths[frame_id + 1] = output
    capture.release()

    success = Image.open(paths[116]).convert("RGB")
    canvas = ImageOps.fit(success, (1600, 900), method=Image.Resampling.LANCZOS)
    canvas = canvas.filter(ImageFilter.GaussianBlur(radius=1.2))
    canvas = ImageEnhance.Contrast(canvas).enhance(1.06)
    overlay = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    pixels = overlay.load()
    for x in range(canvas.width):
        t = x / canvas.width
        alpha = int(215 * max(0.10, 1.0 - 1.25 * t))
        for y in range(canvas.height):
            pixels[x, y] = (7, 10, 20, alpha)
    canvas = Image.alpha_composite(canvas.convert("RGBA"), overlay).convert("RGB")
    cover = ASSET_DIR / "cover_success.jpg"
    canvas.save(cover, quality=94)
    paths[0] = cover
    return paths


def configure_matplotlib() -> None:
    prop = font_manager.FontProperties(fname=FONT_FILE)
    plt.rcParams.update(
        {
            "font.family": prop.get_name(),
            "axes.facecolor": f"#{PANEL}",
            "figure.facecolor": f"#{PANEL}",
            "text.color": f"#{TEXT}",
            "axes.labelcolor": f"#{TEXT}",
            "xtick.color": f"#{MUTED}",
            "ytick.color": f"#{MUTED}",
            "axes.edgecolor": f"#{GRID}",
            "axes.titleweight": "bold",
        }
    )


def make_forward_dynamics_chart(report: dict) -> Path:
    metrics = report["metrics"]
    names = ["Identity", "Delta\nintegration", "Linear", "Ridge\nα=0.01"]
    keys = ["baseline_identity", "baseline_delta_integration", "linear", "ridge_alpha_0.01"]
    positions = [metrics[key]["position_rmse_m"] * 1000 for key in keys]
    orientations = [metrics[key]["orientation_rmse_deg"] for key in keys]
    colors = [BLUE, YELLOW, MUTED, PURPLE]

    fig, axes = plt.subplots(1, 2, figsize=(10.7, 5.85), dpi=180)
    for ax, values, title, unit in [
        (axes[0], positions, "Held-out position RMSE", "mm"),
        (axes[1], orientations, "Held-out orientation RMSE", "deg"),
    ]:
        bars = ax.bar(np.arange(len(names)), values, color=[f"#{c}" for c in colors], width=0.67)
        ax.set_title(title, fontsize=13, color=f"#{WHITE}", pad=10)
        ax.set_xticks(np.arange(len(names)), names, fontsize=9.5)
        ax.grid(axis="y", color=f"#{GRID}", linewidth=0.8, alpha=0.7)
        ax.set_axisbelow(True)
        ax.spines[["top", "right"]].set_visible(False)
        ax.spines[["left", "bottom"]].set_color(f"#{GRID}")
        top = max(values) * 1.28
        ax.set_ylim(0, top)
        for bar, value in zip(bars, values):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                value + top * 0.035,
                f"{value:.2f} {unit}",
                ha="center",
                va="bottom",
                fontsize=9.5,
                color=f"#{WHITE}",
                fontweight="bold",
            )
    fig.tight_layout(pad=1.2, w_pad=2.0)
    output = ASSET_DIR / "forward_dynamics_rmse.png"
    fig.savefig(output, facecolor=f"#{PANEL}")
    plt.close(fig)
    return output


def make_tracking_chart(oracle_report: dict, ridge_report: dict) -> Path:
    oracle = {int(item["seed"]): item for item in oracle_report["reports"]}
    ridge = {int(item["seed"]): item for item in ridge_report["reports"]}
    seeds = [1, 3, 7]
    oracle_pos = [oracle[s]["position_error_mean_m"] * 100 for s in seeds]
    ridge_pos = [ridge[s]["position_error_mean_m"] * 100 for s in seeds]
    oracle_ori = [math.degrees(oracle[s]["orientation_error_mean_rad"]) for s in seeds]
    ridge_ori = [math.degrees(ridge[s]["orientation_error_mean_rad"]) for s in seeds]

    fig, axes = plt.subplots(1, 2, figsize=(10.7, 5.65), dpi=180)
    x = np.arange(len(seeds))
    width = 0.32
    for ax, a, b, title, unit in [
        (axes[0], oracle_pos, ridge_pos, "Mean position tracking error", "cm"),
        (axes[1], oracle_ori, ridge_ori, "Mean orientation tracking error", "deg"),
    ]:
        bars_a = ax.bar(x - width / 2, a, width, label="Oracle", color=f"#{GREEN}")
        bars_b = ax.bar(x + width / 2, b, width, label="Ridge", color=f"#{PURPLE}")
        ax.set_xticks(x, [f"seed {s}" for s in seeds], fontsize=9.5)
        ax.set_title(title, fontsize=13, color=f"#{WHITE}", pad=10)
        ax.grid(axis="y", color=f"#{GRID}", linewidth=0.8, alpha=0.7)
        ax.set_axisbelow(True)
        ax.spines[["top", "right"]].set_visible(False)
        ax.spines[["left", "bottom"]].set_color(f"#{GRID}")
        top = max(a + b) * 1.45
        ax.set_ylim(0, top)
        for bars in (bars_a, bars_b):
            for bar in bars:
                value = bar.get_height()
                ax.text(
                    bar.get_x() + bar.get_width() / 2,
                    value + top * 0.035,
                    f"{value:.2f}",
                    ha="center",
                    va="bottom",
                    fontsize=8.8,
                    color=f"#{WHITE}",
                    fontweight="bold",
                )
        ax.text(0.98, 0.98, unit, transform=ax.transAxes, ha="right", va="top", color=f"#{MUTED}", fontsize=9)
    axes[1].legend(frameon=False, labelcolor=f"#{TEXT}", loc="upper left")
    fig.tight_layout(pad=1.2, w_pad=2.1)
    output = ASSET_DIR / "tracking_oracle_vs_ridge.png"
    fig.savefig(output, facecolor=f"#{PANEL}")
    plt.close(fig)
    return output


def make_lift_success_chart(lift_report: dict) -> Path:
    reports = sorted(lift_report["reports"], key=lambda item: int(item["seed"]))
    seeds = [int(item["seed"]) for item in reports]
    source = [int(item["source_success_step"]) for item in reports]
    target = [int(item["steps"][-1]["step_id"]) + 1 for item in reports]

    fig, ax = plt.subplots(figsize=(8.4, 4.52), dpi=180)
    y = np.arange(len(seeds))
    for idx, (src, tgt) in enumerate(zip(source, target)):
        ax.plot([src, tgt], [idx, idx], color=f"#{GRID}", linewidth=5, solid_capstyle="round")
        ax.scatter(src, idx, s=120, color=f"#{GREEN}", zorder=3)
        ax.scatter(tgt, idx, s=120, color=f"#{ORANGE}", zorder=3)
        ax.text(src - 1.2, idx - 0.24, str(src), color=f"#{GREEN}", fontsize=10, fontweight="bold", ha="right")
        ax.text(tgt + 1.2, idx - 0.24, str(tgt), color=f"#{ORANGE}", fontsize=10, fontweight="bold", ha="left")
    ax.set_yticks(y, [f"seed {seed}" for seed in seeds], fontsize=10)
    ax.set_xlim(35, 62)
    ax.set_xlabel("policy step", fontsize=10)
    ax.invert_yaxis()
    ax.grid(axis="x", color=f"#{GRID}", alpha=0.65)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.spines["bottom"].set_color(f"#{GRID}")
    ax.scatter([], [], s=90, color=f"#{GREEN}", label="Panda source success")
    ax.scatter([], [], s=90, color=f"#{ORANGE}", label="xArm7 target success")
    ax.legend(frameon=False, labelcolor=f"#{TEXT}", loc="lower right", fontsize=9)
    fig.tight_layout(pad=1.1)
    output = ASSET_DIR / "lift_success_steps.png"
    fig.savefig(output, facecolor=f"#{PANEL}")
    plt.close(fig)
    return output


def add_scope_strip(slide, y: float = 6.55) -> None:
    add_rect(slide, 0.64, y, 12.04, 0.47, fill="11192A", line=GRID)
    add_rich_text(
        slide,
        [
            ("범위  ", MUTED, True),
            ("SIMULATION ONLY", ORANGE, True),
            ("  ·  ROS command / xArm SDK / robot IP / real calibration 사용 안 함", TEXT, False),
        ],
        0.92,
        y + 0.04,
        11.45,
        0.36,
        size=12.5,
    )


def build_presentation() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    configure_matplotlib()
    frames = make_video_assets()
    fd_report = load_json(FD_REPORT_PATH)
    oracle_report = load_json(TRACK_ORACLE_PATH)
    ridge_report = load_json(TRACK_RIDGE_PATH)
    pick_report = load_json(PICK_REPORT_PATH)
    lift_report = load_json(LIFT_REPORT_PATH)
    fd_chart = make_forward_dynamics_chart(fd_report)
    tracking_chart = make_tracking_chart(oracle_report, ridge_report)
    lift_chart = make_lift_success_chart(lift_report)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "Panda to xArm7 Simulation Policy Transfer"
    prs.core_properties.subject = "Current Mirage xArm7 experiment setup and validation"
    prs.core_properties.author = "Mirage experiment deck"
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    add_picture_crop(slide, frames[0], 0, 0, 13.333, 7.5)
    add_badge(slide, "MIRAGE · CROSS-EMBODIMENT", 0.72, 0.66, 2.90, fill="142134", color=ORANGE, line="2C3A51")
    add_text(slide, "Panda → xArm7", 0.72, 1.58, 6.55, 0.68, size=32, color=WHITE, bold=True)
    add_text(slide, "Simulation Policy Transfer", 0.72, 2.25, 6.70, 0.80, size=35, color=WHITE, bold=True)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.74), Inches(3.20), Inches(0.74), Inches(0.06))
    set_fill(accent, ORANGE)
    accent.line.fill.background()
    add_text(
        slide,
        "현재 실험 세팅 · 전이 파이프라인 · 정량 검증",
        0.72,
        3.48,
        5.85,
        0.48,
        size=18,
        color=TEXT,
        bold=True,
    )
    add_text(
        slide,
        "PickPlaceCan + Lift  |  Robosuite 1.4.1  |  low-dimensional policy",
        0.72,
        4.08,
        6.30,
        0.36,
        size=12.5,
        color=MUTED,
    )
    add_rect(slide, 0.72, 5.61, 4.15, 0.87, fill="10182A", line=GRID)
    add_rich_text(
        slide,
        [
            ("TARGET SUCCESS  ", GREEN, True),
            ("PickPlaceCan step 116", WHITE, True),
        ],
        0.98,
        5.78,
        3.68,
        0.28,
        size=14,
    )
    add_text(slide, "2026.08.10 · local working tree", 0.98, 6.12, 3.50, 0.23, size=9.5, color=MUTED)

    # 2. Executive outcome
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "한 줄 결론", "정책 실행 주체와 성공 판정 주체를 분리해 xArm7 전이를 검증", 2)
    add_text(
        slide,
        "정책은 Panda에서 실행하고,\n성공은 xArm7에서 판정했다.",
        0.74,
        1.64,
        5.05,
        1.24,
        size=28,
        color=WHITE,
        bold=True,
    )
    add_text(
        slide,
        "동일한 저차원 체크포인트 출력을 Panda 시뮬레이터에서 1 step 실행한 뒤,\n달성 pose를 변환·검증하여 xArm7이 절대 pose로 추종한다.",
        0.76,
        3.03,
        5.06,
        0.95,
        size=14.5,
        color=TEXT,
    )
    add_rect(slide, 0.76, 4.42, 5.05, 1.26, fill="111B2B", line=GRID)
    add_rich_text(
        slide,
        [("의미  ", ORANGE, True), ("로봇 형태가 달라도 task-space pose contract로 실행 경로를 분리", WHITE, True)],
        1.02,
        4.63,
        4.55,
        0.42,
        size=13.2,
    )
    add_text(slide, "현재 결과는 실제 로봇이 아닌 시뮬레이션 검증이다.", 1.02, 5.12, 4.52, 0.28, size=10.5, color=MUTED)
    add_metric_card(slide, 6.20, 1.60, 2.05, 3.65, "1 / 1", "PickPlaceCan", "target step 116\ncontact-relaxed steps 32", accent=GREEN)
    add_metric_card(slide, 8.48, 1.60, 2.05, 3.65, "4 / 4", "Lift", "seeds 0, 1, 3, 7\ntarget step 41–57", accent=ORANGE)
    add_metric_card(slide, 10.76, 1.60, 1.92, 3.65, "0", "실패 이벤트", "timeout 0\niteration limit 0\nrejection 0", accent=BLUE)
    add_rect(slide, 6.20, 5.55, 6.48, 0.72, fill=GREEN_DARK, line="2A725A")
    add_rich_text(
        slide,
        [("검증  ", GREEN, True), ("현재 작업 트리 회귀 테스트 87 passed", WHITE, True)],
        6.53,
        5.72,
        5.85,
        0.34,
        size=14.5,
    )
    add_footer(slide, "sources: task-aware rollout reports + pytest")

    # 3. Experiment setup
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "실험 세팅", "source policy와 target control environment를 완전히 분리", 3)
    add_rect(slide, 0.70, 1.60, 5.20, 4.58, fill=PANEL, line="2A4350")
    add_badge(slide, "SOURCE", 0.98, 1.86, 1.10, fill=GREEN_DARK, color=GREEN, line="2B6B56")
    add_text(slide, "Franka Emika Panda", 0.98, 2.34, 4.35, 0.45, size=22, color=WHITE, bold=True)
    add_text(slide, "Robosuite robot id: Panda", 0.98, 2.82, 4.35, 0.30, size=11, color=MUTED)
    add_multiline(
        slide,
        [
            ("저차원 robomimic checkpoint", TEXT, False),
            ("action: [Δp(3), Δrotvec(3), gripper]", TEXT, False),
            ("OSC_POSE · control_delta=true · 20 Hz", TEXT, False),
            ("checkpoint action + 1 physics step", TEXT, False),
            ("camera / offscreen rendering disabled", TEXT, False),
        ],
        1.02,
        3.38,
        4.35,
        2.18,
        size=14,
        bullet=True,
        spacing_after=8,
    )
    add_badge(slide, "panda_sim_world", 1.00, 5.64, 1.73, fill="18283A", color=GREEN, size=10)

    add_rect(slide, 7.43, 1.60, 5.20, 4.58, fill=PANEL, line="4B372A")
    add_badge(slide, "TARGET", 7.72, 1.86, 1.12, fill=ORANGE_DARK, color=ORANGE, line="7E4C30")
    add_text(slide, "xArm7", 7.72, 2.34, 4.35, 0.45, size=22, color=WHITE, bold=True)
    add_text(slide, "in-process Robosuite target", 7.72, 2.82, 4.35, 0.30, size=11, color=MUTED)
    add_multiline(
        slide,
        [
            ("policy camera를 상속하지 않는 control env", TEXT, False),
            ("absolute blocking OSC_POSE", TEXT, False),
            ("target task success로만 early stop", TEXT, False),
            ("workspace / joint / freshness safety gate", TEXT, False),
            ("gripper: −1 open, +1 closed", TEXT, False),
        ],
        7.76,
        3.38,
        4.35,
        2.18,
        size=14,
        bullet=True,
        spacing_after=8,
    )
    add_badge(slide, "xarm7_sim_world", 7.74, 5.64, 1.76, fill="33271F", color=ORANGE, size=10)
    add_arrow(slide, 5.94, 3.56, 7.30, 3.56, color=ORANGE, width=2.6)
    add_text(slide, "pose contract", 5.94, 3.08, 1.33, 0.28, size=10.5, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_scope_strip(slide)
    add_footer(slide, "source: xarm_transfer_sim.example.yaml")

    # 4. Pipeline
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "End-to-end 실행 파이프라인", "온라인 oracle과 learned Ridge가 같은 계약·안전 계층을 공유", 4)
    stages = [
        ("01", "Checkpoint", "raw action\n+ metadata", GREEN),
        ("02", "Panda step", "one policy step\n+ achieved pose", GREEN),
        ("03", "Transform v2", "task world ×\n tool frame", BLUE),
        ("04", "Dynamics", "online oracle\n or Ridge", PURPLE),
        ("05", "Safety gate", "fresh · finite\n workspace · joints", YELLOW),
        ("06", "xArm7 OSC", "absolute blocking\n + outer loop", ORANGE),
        ("07", "Target success", "task event\n terminates", GREEN),
    ]
    start_x = 0.72
    y = 2.14
    box_w = 1.54
    gap = 0.21
    for idx, (num, name, desc, accent) in enumerate(stages):
        x = start_x + idx * (box_w + gap)
        add_rect(slide, x, y, box_w, 2.52, fill=PANEL, line=GRID)
        add_text(slide, num, x + 0.17, y + 0.16, 0.45, 0.30, size=10.5, color=accent, bold=True)
        add_text(slide, name, x + 0.17, y + 0.62, box_w - 0.30, 0.52, size=15, color=WHITE, bold=True)
        add_text(slide, desc, x + 0.17, y + 1.34, box_w - 0.30, 0.72, size=11.2, color=MUTED)
        if idx < len(stages) - 1:
            add_arrow(slide, x + box_w + 0.03, y + 1.26, x + box_w + gap - 0.05, y + 1.26, color="6D7C93", width=1.6)
    add_rect(slide, 2.46, 5.10, 6.84, 0.92, fill="121A2B", line=GRID)
    add_rich_text(
        slide,
        [
            ("Provider만 교체  ", PURPLE, True),
            ("oracle: Panda achieved next pose", TEXT, False),
            ("   |   ", MUTED, False),
            ("Ridge: autoregressive predicted pose", TEXT, False),
        ],
        2.76,
        5.34,
        6.26,
        0.36,
        size=13,
        align=PP_ALIGN.CENTER,
    )
    add_arrow(slide, 9.90, 5.50, 10.94, 4.76, color=ORANGE, width=1.6)
    add_text(slide, "achieved pose feedback", 9.22, 5.58, 2.44, 0.30, size=10, color=MUTED, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.72, 6.29, 11.96, 0.50, fill=GREEN_DARK, line="2A715A")
    add_text(slide, "같은 frame·unit·timestamp contract와 fail-closed validation을 모든 단계에서 유지", 0.98, 6.39, 11.44, 0.28, size=13.2, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "source: rollout.py · contracts.py · runtime.py")

    # 5. Transform
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "좌표 변환 v2", "task/world 정합과 tool convention을 분리해 motion direction 왜곡을 방지", 5)
    add_rect(slide, 0.72, 1.62, 11.92, 0.98, fill="111A2C", line=GRID)
    add_rich_text(
        slide,
        [
            ("T", MUTED, False),
            ("targetWorld,targetTool", ORANGE, True),
            ("  =  T", MUTED, False),
            ("targetWorld,sourceWorld", BLUE, True),
            ("  @  T", MUTED, False),
            ("sourceWorld,sourceTool", GREEN, True),
            ("  @  T", MUTED, False),
            ("sourceTool,targetTool", PURPLE, True),
        ],
        1.08,
        1.88,
        11.18,
        0.46,
        size=18,
        align=PP_ALIGN.CENTER,
    )
    add_rect(slide, 0.74, 2.94, 5.58, 2.74, fill=PANEL, line="28516D")
    add_badge(slide, "LEFT-MULTIPLY", 1.02, 3.20, 1.66, fill=BLUE_DARK, color=BLUE, line="356A88")
    add_text(slide, "Task / world transform", 1.02, 3.70, 4.82, 0.44, size=20, color=WHITE, bold=True)
    add_text(slide, "Panda와 xArm7이 동일한 task 좌표를 사용", 1.02, 4.20, 4.82, 0.36, size=13.5, color=TEXT)
    add_rich_text(slide, [("현재 값  ", MUTED, True), ("IDENTITY", BLUE, True)], 1.02, 4.72, 4.70, 0.32, size=14)
    add_text(slide, "translation = [0, 0, 0] m\nrotation = [0, 0, 0, 1] XYZW", 1.02, 5.08, 4.60, 0.54, size=11.5, color=MUTED, font=MONO)
    add_rect(slide, 6.58, 2.94, 6.06, 2.74, fill=PANEL, line="554D8D")
    add_badge(slide, "RIGHT-MULTIPLY", 6.88, 3.20, 1.78, fill=PURPLE_DARK, color=PURPLE, line="655BB0")
    add_text(slide, "Tool-frame correction", 6.88, 3.70, 5.12, 0.44, size=20, color=WHITE, bold=True)
    add_text(slide, "nominal gripper axis 차이만 local-Z에서 보정", 6.88, 4.20, 5.12, 0.36, size=13.5, color=TEXT)
    add_rich_text(slide, [("현재 값  ", MUTED, True), ("−90° about local Z", PURPLE, True)], 6.88, 4.72, 4.95, 0.32, size=14)
    add_text(slide, "quaternion = [0, 0, −0.7071, 0.7071]", 6.88, 5.18, 5.08, 0.30, size=11.5, color=MUTED, font=MONO)
    add_rect(slide, 0.74, 6.03, 11.90, 0.67, fill="182233", line=GRID)
    add_rich_text(
        slide,
        [
            ("핵심  ", ORANGE, True),
            ("home pose로 변환을 추정하지 않으며, simulation transform을 real calibration으로 표시할 수 없다.", WHITE, True),
        ],
        1.02,
        6.18,
        11.36,
        0.36,
        size=13.2,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "source: simulation_transform.py · transform version 2")

    # 6. Control and safety
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "Target control & safety", "absolute blocking OSC에 보상 루프와 task-aware 예외를 제한적으로 적용", 6)
    add_rect(slide, 0.72, 1.64, 4.01, 4.93, fill=PANEL, line=GRID)
    add_badge(slide, "ADMISSION GATES", 1.02, 1.92, 1.86, fill=BLUE_DARK, color=BLUE, line="345E7A")
    gates = [
        ("Frame / type", "simulation_only · xarm7_sim_world"),
        ("Freshness", "age ≤ 0.15 s · future skew ≤ 0.02 s"),
        ("Workspace", "[−0.65,−0.60,0.40]–[0.35,0.60,1.35] m"),
        ("Joint limits", "model limit + 0.01 rad margin"),
        ("Learned model", "held-out position / orientation bounds"),
    ]
    y = 2.52
    for name, detail in gates:
        add_text(slide, name, 1.02, y, 1.30, 0.32, size=12.3, color=WHITE, bold=True)
        add_text(slide, detail, 2.29, y, 2.09, 0.40, size=9.5, color=MUTED)
        y += 0.70
    add_rect(slide, 4.97, 1.64, 3.62, 4.93, fill=PANEL, line=GRID)
    add_badge(slide, "BLOCKING OSC", 5.26, 1.92, 1.66, fill=ORANGE_DARK, color=ORANGE, line="7D4A2D")
    add_text(slide, "Free-space convergence", 5.26, 2.55, 2.88, 0.36, size=16.5, color=WHITE, bold=True)
    add_metric_card(slide, 5.24, 3.04, 1.39, 1.55, "2.5", "cm", "position tol.", accent=ORANGE)
    add_metric_card(slide, 6.82, 3.04, 1.39, 1.55, "7°", "orientation", "tolerance", accent=ORANGE)
    add_multiline(
        slide,
        [
            ("5 consecutive in-bound steps", TEXT, False),
            ("position correction ≤ 3 cm", TEXT, False),
            ("orientation correction ≤ 5°", TEXT, False),
            ("timeout 2 s · max 300 iterations", TEXT, False),
        ],
        5.28,
        4.82,
        2.86,
        1.52,
        size=11,
        bullet=True,
        spacing_after=4,
    )
    add_rect(slide, 8.83, 1.64, 3.81, 4.93, fill=PANEL, line="5B4335")
    add_badge(slide, "CONTACT EXCEPTION", 9.12, 1.92, 2.11, fill=ORANGE_DARK, color=ORANGE, line="7D4A2D")
    add_text(slide, "PickPlaceCan only", 9.12, 2.55, 3.02, 0.36, size=16.5, color=WHITE, bold=True)
    add_text(slide, "simulator가 gripper ↔ Can 접촉을 직접 확인한 동안만", 9.12, 2.99, 3.05, 0.55, size=11.5, color=TEXT)
    add_rich_text(slide, [("6.0 cm", ORANGE, True), ("  /  ", MUTED, False), ("18°", ORANGE, True)], 9.12, 3.78, 3.04, 0.42, size=21, align=PP_ALIGN.CENTER)
    add_multiline(
        slide,
        [
            ("Kp = [300,300,500,300,300,300]", TEXT, False),
            ("position compensation ≤ 10 cm", TEXT, False),
            ("Lift에서는 항상 비활성", TEXT, False),
            ("workspace/joint/freshness gate는 유지", TEXT, False),
        ],
        9.13,
        4.50,
        3.06,
        1.58,
        size=10.8,
        bullet=True,
        spacing_after=4,
    )
    add_rect(slide, 0.72, 6.72, 11.92, 0.30, fill="2B2030", line=None)
    add_text(slide, "Target task-success event가 발생하면 마지막 EE pose가 미수렴이어도 즉시 종료", 0.96, 6.745, 11.45, 0.24, size=10.8, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "source: profiles.py · control.py")

    # 7. Evaluation protocol
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "평가 프로토콜", "task profile이 checkpoint·observation·controller를 정적으로 고정", 7)
    add_rect(slide, 0.72, 1.60, 5.74, 2.25, fill=PANEL, line="2A664F")
    add_badge(slide, "TASK 01", 0.98, 1.86, 1.05, fill=GREEN_DARK, color=GREEN, line="2B6B56")
    add_text(slide, "PickPlaceCan", 0.98, 2.34, 2.80, 0.40, size=21, color=WHITE, bold=True)
    add_multiline(
        slide,
        [
            ("object obs [14] · graspable Can", TEXT, False),
            ("checkpoint: can / low_dim / epoch 250", TEXT, False),
            ("native horizon 400 · full task seed 1", TEXT, False),
            ("contact relaxation enabled", ORANGE, True),
        ],
        3.12,
        1.92,
        3.00,
        1.55,
        size=11.5,
        bullet=True,
        spacing_after=5,
    )
    add_rect(slide, 6.70, 1.60, 5.94, 2.25, fill=PANEL, line="6B4B30")
    add_badge(slide, "TASK 02", 6.98, 1.86, 1.05, fill=ORANGE_DARK, color=ORANGE, line="7D4A2D")
    add_text(slide, "Lift", 6.98, 2.34, 2.05, 0.40, size=21, color=WHITE, bold=True)
    add_multiline(
        slide,
        [
            ("object obs [10] · graspable cube", TEXT, False),
            ("checkpoint: lift / low_dim / epoch 1000", TEXT, False),
            ("native horizon 400 · seeds 0,1,3,7", TEXT, False),
            ("contact relaxation disabled", BLUE, True),
        ],
        9.04,
        1.92,
        3.18,
        1.55,
        size=11.5,
        bullet=True,
        spacing_after=5,
    )
    add_rect(slide, 0.72, 4.15, 4.03, 2.28, fill=PANEL, line=GRID)
    add_badge(slide, "PROVIDERS", 0.98, 4.42, 1.34, fill=PURPLE_DARK, color=PURPLE, line="655BB0")
    add_multiline(
        slide,
        [
            ("online_oracle", GREEN, True),
            ("Panda achieved next pose를 즉시 사용", TEXT, False),
            ("learned_ridge", PURPLE, True),
            ("recorded action을 autoregressive replay", TEXT, False),
        ],
        1.00,
        4.94,
        3.38,
        1.25,
        size=11.4,
        spacing_after=4,
    )
    add_rect(slide, 4.99, 4.15, 4.02, 2.28, fill=PANEL, line=GRID)
    add_badge(slide, "DATASET", 5.26, 4.42, 1.18, fill=BLUE_DARK, color=BLUE, line="345E7A")
    add_rich_text(slide, [("300", WHITE, True), (" transitions", MUTED, False)], 5.28, 4.93, 2.80, 0.44, size=22)
    add_text(slide, "train: seeds 0,2,4,5 (200)\nvalidation: seeds 1,3 (100)\nwhole-episode split · seed 11", 5.28, 5.42, 3.30, 0.82, size=11.2, color=TEXT)
    add_rect(slide, 9.25, 4.15, 3.39, 2.28, fill=PANEL, line=GRID)
    add_badge(slide, "RUNTIME CONTRACT", 9.52, 4.42, 1.92, fill="2A243B", color=YELLOW, line="64567B")
    add_multiline(
        slide,
        [
            ("low_dim만 실행 가능", WHITE, True),
            ("image modes: classified, fail closed", TEXT, False),
            ("unregistered task: fail closed", TEXT, False),
            ("Lift Ridge training: 현재 범위 밖", MUTED, False),
        ],
        9.54,
        4.95,
        2.70,
        1.26,
        size=11,
        spacing_after=4,
    )
    add_footer(slide, "source: profiles.py · dataset/model reports")

    # 8. Forward dynamics
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "Forward dynamics: held-out one-step prediction", "episode leakage 없이 100개 validation transition으로 비교", 8)
    add_rect(slide, 0.72, 1.55, 8.56, 4.92, fill=PANEL, line=GRID)
    add_picture_crop(slide, fd_chart, 0.93, 1.78, 8.14, 4.45)
    add_metric_card(slide, 9.55, 1.55, 3.08, 1.48, "2.39 mm", "Ridge position RMSE", "α = 0.01 · selected deployable model", accent=PURPLE)
    add_metric_card(slide, 9.55, 3.25, 3.08, 1.48, "0.254°", "Ridge orientation RMSE", "held-out one-step prediction", accent=PURPLE)
    add_rect(slide, 9.55, 4.95, 3.08, 1.52, fill=GREEN_DARK, line="2B6B56")
    add_rich_text(slide, [("≈75%↓", GREEN, True), ("  position RMSE", WHITE, True)], 9.81, 5.18, 2.54, 0.37, size=19, align=PP_ALIGN.CENTER)
    add_text(slide, "Ridge vs. identity baseline\n9.63 mm → 2.39 mm", 9.82, 5.70, 2.52, 0.48, size=11, color=TEXT, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.72, 6.69, 11.91, 0.31, fill="151E31", line=None)
    add_text(slide, "주의: one-step held-out error는 50-step autoregressive drift의 online safety bound가 아니다.", 0.94, 6.72, 11.48, 0.24, size=10.5, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "source: forward_dynamics_evaluation.json")

    # 9. Tracking comparison
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "50-step xArm7 tracking: Oracle vs. Ridge", "동일 held-out seeds 1·3·7, 같은 transform·controller·safety stack", 9)
    add_rect(slide, 0.72, 1.55, 8.70, 4.86, fill=PANEL, line=GRID)
    add_picture_crop(slide, tracking_chart, 0.92, 1.80, 8.30, 4.38)
    add_metric_card(slide, 9.69, 1.55, 2.95, 1.48, "6 / 6", "rollouts complete", "Oracle 3 + Ridge 3\nall 50 / 50 steps", accent=GREEN)
    add_metric_card(slide, 9.69, 3.25, 2.95, 1.48, "0", "control failures", "timeout · iteration limit\nrejection · contact event", accent=BLUE)
    add_rect(slide, 9.69, 4.95, 2.95, 1.46, fill=PURPLE_DARK, line="655BB0")
    add_text(slide, "Ridge도 target tracking 유지", 9.94, 5.17, 2.44, 0.42, size=14.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "mean pos 0.49–0.89 cm\nmean ori 1.02–2.17°", 9.94, 5.72, 2.44, 0.48, size=11.2, color=TEXT, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.72, 6.66, 11.91, 0.35, fill="172536", line=None)
    add_text(slide, "그래프는 controller가 달성한 target tracking mean error; learned source-pose drift와 분리해 보고", 0.94, 6.70, 11.48, 0.25, size=10.6, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "source: task_tool_v2 oracle / learned rollout reports")

    # 10. PickPlaceCan full task
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "Full-task 결과 ① PickPlaceCan", "xArm7가 can을 집어 bin에 놓는 실제 task event로 성공 판정", 10)
    image_specs = [
        (frames[1], 0.72, 1.60, "STEP 1 · approach"),
        (frames[31], 5.05, 1.60, "STEP 31 · pre-grasp"),
        (frames[61], 0.72, 4.02, "STEP 61 · grasp"),
        (frames[116], 5.05, 4.02, "STEP 116 · TARGET SUCCESS"),
    ]
    for path, x, y, label in image_specs:
        add_picture_crop(slide, path, x, y, 4.02, 2.02, border=GRID)
        fill = GREEN_DARK if "SUCCESS" in label else "101827"
        add_rect(slide, x + 0.12, y + 1.58, 2.56 if "SUCCESS" in label else 1.92, 0.31, fill=fill, line=None)
        add_text(slide, label, x + 0.20, y + 1.60, 2.40, 0.24, size=9.2, color=GREEN if "SUCCESS" in label else WHITE, bold=True)
    add_rect(slide, 9.34, 1.60, 3.29, 4.44, fill=PANEL, line=GRID)
    add_badge(slide, "SEED 1", 9.63, 1.88, 1.02, fill=GREEN_DARK, color=GREEN, line="2B6B56")
    add_text(slide, "116 / 400", 9.63, 2.42, 2.44, 0.63, size=29, color=GREEN, bold=True)
    add_text(slide, "policy steps until target success", 9.64, 3.03, 2.48, 0.36, size=10.5, color=MUTED)
    add_multiline(
        slide,
        [
            ("timeout 0 · iteration limit 0", TEXT, False),
            ("safety rejection 0", TEXT, False),
            ("verified-contact steps 32", ORANGE, True),
            ("mean command latency 38.6 ms", TEXT, False),
            ("target-only early termination", TEXT, False),
        ],
        9.65,
        3.65,
        2.46,
        1.68,
        size=11.2,
        bullet=True,
        spacing_after=6,
    )
    add_rect(slide, 9.62, 5.45, 2.71, 0.39, fill="2B2030", line=None)
    add_text(slide, "final task event > pose convergence", 9.74, 5.52, 2.46, 0.24, size=9.2, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.72, 6.35, 11.91, 0.50, fill=GREEN_DARK, line="2A715A")
    add_text(slide, "성공 시점의 큰 EE residual은 object contact / immediate task-success 규칙으로 별도 기록됨", 0.96, 6.46, 11.45, 0.28, size=11.6, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "source: panda_xarm_oracle_seed1_full_task.mp4 + report")

    # 11. Lift full task
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "Full-task 결과 ② Lift", "모든 seed에서 target success까지 계속 실행; contact relaxation은 사용하지 않음", 11)
    add_rect(slide, 0.72, 1.55, 8.56, 4.97, fill=PANEL, line=GRID)
    add_picture_crop(slide, lift_chart, 0.98, 1.86, 8.06, 4.34)
    add_metric_card(slide, 9.55, 1.55, 3.08, 1.48, "4 / 4", "target success", "seeds 0 · 1 · 3 · 7", accent=ORANGE)
    add_metric_card(slide, 9.55, 3.25, 3.08, 1.48, "≤ 2.07 cm", "maximum position error", "all four full-task rollouts", accent=BLUE)
    add_rect(slide, 9.55, 4.95, 3.08, 1.57, fill=GREEN_DARK, line="2B6B56")
    add_text(slide, "0 failures", 9.82, 5.16, 2.54, 0.38, size=21, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "timeout · iteration limit\nrejection · contact tolerance", 9.81, 5.68, 2.56, 0.58, size=10.8, color=TEXT, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.72, 6.71, 11.91, 0.31, fill="151E31", line=None)
    add_text(slide, "seed 1: Panda success step 40 → xArm7 success step 41; source success가 target rollout을 조기 종료하지 않음", 0.94, 6.735, 11.48, 0.24, size=10.5, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "source: lift/low_dim/xarm_rollout_report.json")

    # 12. Reproducibility
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "재현 방법", "task-aware CLI가 checkpoint에서 task·mode·native horizon을 해석", 12)
    add_rect(slide, 0.72, 1.56, 7.34, 4.97, fill="101726", line=GRID)
    add_badge(slide, "RUN", 1.02, 1.84, 0.72, fill=ORANGE_DARK, color=ORANGE, line="7D4A2D")
    command = (
        "cd mirage\n"
        "MUJOCO_GL=egl ../.venv/bin/python -m \\\n"
        "  mirage.benchmark.robosuite.xarm_transfer.cli rollout \\\n"
        "  --checkpoint mirage/models/can/low_dim/\\\n"
        "    model_epoch_250_PickPlaceCan_success_1.0.pth \\\n"
        "  --mode oracle --seeds 1 --terminate-on-success"
    )
    add_text(slide, command, 1.02, 2.36, 6.72, 2.35, size=11.1, color="D7E2EF", font=MONO)
    add_badge(slide, "VERIFY", 1.02, 4.86, 0.98, fill=GREEN_DARK, color=GREEN, line="2B6B56")
    verify = (
        "PYTEST_DISABLE_PLUGIN_AUTOLOAD=1 PYTHONPATH=... \\\n"
        "  python -m pytest -q --import-mode=importlib \\\n"
        "  test_xarm_transfer_*.py test_canonical_policy_checkpoint.py"
    )
    add_text(slide, verify, 1.02, 5.30, 6.70, 0.88, size=9.5, color="D7E2EF", font=MONO)
    add_rect(slide, 8.33, 1.56, 4.30, 2.68, fill=PANEL, line=GRID)
    add_badge(slide, "ENVIRONMENT", 8.61, 1.84, 1.48, fill=BLUE_DARK, color=BLUE, line="345E7A")
    add_multiline(
        slide,
        [
            ("Ubuntu 22.04.5 · Python 3.10.12", TEXT, False),
            ("Robosuite 1.4.1 · MuJoCo 2.3.2", TEXT, False),
            ("ROS 2 Humble installed, unused in rollout", TEXT, False),
            ("RTX 5060 Laptop GPU · EGL rendering", TEXT, False),
            ("branch: agent/xarm7-development", TEXT, False),
            ("HEAD: 4e4429e + local working tree", MUTED, False),
        ],
        8.62,
        2.40,
        3.55,
        1.58,
        size=11.1,
        bullet=True,
        spacing_after=5,
    )
    add_rect(slide, 8.33, 4.51, 4.30, 2.02, fill=PANEL, line=GRID)
    add_badge(slide, "VALIDATION", 8.61, 4.79, 1.28, fill=GREEN_DARK, color=GREEN, line="2B6B56")
    add_rich_text(slide, [("87 passed", GREEN, True), ("  in 70.81 s", MUTED, False)], 8.62, 5.34, 3.45, 0.40, size=22)
    add_text(slide, "current working tree · 2026-08-10\nresults stay under ignored results/", 8.63, 5.90, 3.43, 0.48, size=10.5, color=TEXT)
    add_footer(slide, "source: CLI README + local validation")

    # 13. Closing
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, fill=BG, radius=False)
    add_title(slide, "결론 및 다음 단계", "현재 성과와 경계 조건을 함께 제시", 13)
    add_text(slide, "현재 확보한 것", 0.76, 1.66, 3.30, 0.42, size=21, color=GREEN, bold=True)
    add_multiline(
        slide,
        [
            ("Panda policy → xArm7 control의 명시적 contract", WHITE, True),
            ("task/world와 tool-frame을 분리한 transform v2", WHITE, True),
            ("PickPlaceCan 1/1, Lift 4/4 target success", WHITE, True),
            ("Oracle/Ridge 50-step 6/6 completion", WHITE, True),
            ("87개 회귀·통합 테스트 통과", WHITE, True),
        ],
        0.78,
        2.27,
        4.02,
        2.85,
        size=14.2,
        bullet=True,
        spacing_after=13,
    )
    add_text(slide, "아직 주장하지 않는 것", 4.95, 1.66, 3.65, 0.42, size=21, color=YELLOW, bold=True)
    add_multiline(
        slide,
        [
            ("실제 xArm7 hardware 성공", TEXT, False),
            ("real-world calibration", TEXT, False),
            ("image observation runtime", TEXT, False),
            ("Lift learned Ridge model", TEXT, False),
            ("Stack/Square/Assembly 지원", TEXT, False),
        ],
        4.97,
        2.27,
        3.31,
        2.85,
        size=14.2,
        bullet=True,
        spacing_after=13,
    )
    add_rect(slide, 8.55, 1.62, 4.08, 4.70, fill=PANEL, line="4B3F32")
    add_badge(slide, "NEXT", 8.86, 1.92, 0.78, fill=ORANGE_DARK, color=ORANGE, line="7D4A2D")
    add_multiline(
        slide,
        [
            ("1. real calibration 경로를 별도 구현", WHITE, True),
            ("2. low-speed fake → real safety gate", WHITE, True),
            ("3. image observation / cross-painting 연결", WHITE, True),
            ("4. Lift용 dynamics dataset·Ridge 학습", WHITE, True),
            ("5. multi-seed PickPlaceCan 확대", WHITE, True),
        ],
        8.86,
        2.55,
        3.23,
        2.83,
        size=13.6,
        spacing_after=13,
    )
    add_rect(slide, 0.74, 6.54, 11.89, 0.51, fill=GREEN_DARK, line="2A715A")
    add_text(slide, "Takeaway  ·  embodiment 차이는 pose contract로 흡수하고, 안전·성공 판정은 target에서 닫는다.", 0.98, 6.65, 11.42, 0.28, size=13.6, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "Mirage xArm7 simulation transfer")

    prs.save(PPTX_PATH)
    print(PPTX_PATH)


if __name__ == "__main__":
    build_presentation()
